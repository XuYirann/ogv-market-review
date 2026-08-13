"""Extract audience matrices used by the website without modifying the source workbook."""

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data_for_HTML" / "platform_audience.xlsx"
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "app" / "data"
METRICS = {
    "DAU": ("platform_audience_DAU", "platformAudienceDau.json"),
    "MAU": ("platform_audience_MAU", "platformAudienceMau.json"),
}
TOTAL_DURATION_SHEET = "platform_audience_总时长"
TOTAL_DURATION_OUTPUT = "platformAudienceTotalDuration.json"
OVERLAP_SHEET = "红果vs爱腾MAU构成对比_亿"
OVERLAP_OUTPUT = "platformAudienceOverlap.json"


def main() -> None:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    for metric, (sheet_name, filename) in METRICS.items():
        sheet = workbook[sheet_name]
        headers = [sheet.cell(1, column).value for column in range(2, sheet.max_column + 1)]
        valid_columns = [index + 2 for index, value in enumerate(headers) if value is not None]
        periods = [str(sheet.cell(1, column).value)[:7] for column in valid_columns]

        series = []
        for row in range(2, sheet.max_row + 1):
            platform = sheet.cell(row, 1).value
            if not platform:
                continue
            values = []
            for column in valid_columns:
                value = sheet.cell(row, column).value
                values.append(None if value is None else round(float(value) / 10000, 4))
            series.append({"name": str(platform), "values": values})

        payload = {
            "metric": metric,
            "unit": "亿",
            "source": "QuestMobile",
            "periods": periods,
            "series": series,
        }
        output = OUTPUT_DIR / filename
        output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        print(f"Extracted {len(series)} series x {len(periods)} months to {output}")

    sheet = workbook[TOTAL_DURATION_SHEET]
    rows = list(sheet.iter_rows(min_row=1, max_row=10, max_col=103, values_only=True))
    months = [str(value)[:7] for value in rows[0][1:103]]
    quarters = []
    quarter_columns = []
    for index in range(0, len(months), 3):
        quarter_months = months[index:index + 3]
        if len(quarter_months) < 3:
            continue
        year, month = map(int, quarter_months[-1].split("-"))
        quarters.append(f"{str(year)[2:]}Q{(month - 1) // 3 + 1}")
        quarter_columns.append(list(range(index, index + 3)))

    name_map = {
        "长视频平台汇总": "其他长视频平台",
        "短剧平台汇总": "其他短剧平台",
    }
    duration_series = []
    for row in rows[1:]:
        raw_name = row[0]
        if not raw_name:
            continue
        values = []
        for columns in quarter_columns:
            raw_values = [row[column + 1] for column in columns]
            if all(value is None for value in raw_values):
                values.append(None)
            else:
                # Source values are monthly 万分钟; quarterly total is converted to 亿小时.
                values.append(round(sum(float(value or 0) for value in raw_values) / 600000, 4))
        duration_series.append({"name": name_map.get(str(raw_name), str(raw_name)), "values": values})

    duration_payload = {
        "metric": "总用户时长",
        "unit": "亿小时",
        "source": "QuestMobile",
        "periods": quarters,
        "series": duration_series,
    }
    duration_output = OUTPUT_DIR / TOTAL_DURATION_OUTPUT
    duration_output.write_text(json.dumps(duration_payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Extracted {len(duration_series)} series x {len(quarters)} quarters to {duration_output}")

    overlap_sheet = workbook[OVERLAP_SHEET]
    overlap_rows = list(overlap_sheet.iter_rows(min_row=1, max_row=4, max_col=5, values_only=True))
    overlap_periods = []
    for value in overlap_rows[0][1:]:
        if isinstance(value, datetime):
            overlap_periods.append(value.strftime("%Y-%m"))
        elif isinstance(value, (int, float)):
            overlap_periods.append(from_excel(value, workbook.epoch).strftime("%Y-%m"))
        else:
            overlap_periods.append(str(value).replace("年", "-").replace("月", "").replace(" ", ""))
    overlap_payload = {
        "metric": "MAU构成",
        "unit": "亿",
        "source": "QuestMobile",
        "periods": overlap_periods,
        "series": [
            {"name": str(row[0]), "values": [round(float(value), 6) for value in row[1:]]}
            for row in overlap_rows[1:]
        ],
    }
    overlap_output = OUTPUT_DIR / OVERLAP_OUTPUT
    overlap_output.write_text(json.dumps(overlap_payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Extracted overlap composition to {overlap_output}")


if __name__ == "__main__":
    main()
