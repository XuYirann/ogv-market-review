"""Extract horizontal-drama quarterly totals used by the website."""

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data_for_HTML" / "分品类_横屏剧集.xlsx"
OUTPUT = Path(__file__).resolve().parents[1] / "app" / "data" / "horizontalDramaTrend.json"


def read_sheet(workbook, name):
    sheet = workbook[name]
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not values[0]:
            continue
        row = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
        rows.append(row)
    return rows


def clean(value):
    if value is None:
        return None
    number = float(value)
    return 0 if abs(number) < 1e-9 else round(number, 4)


def main():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    totals = read_sheet(workbook, "byQ_长剧有效播放_亿")
    platforms = read_sheet(workbook, "byQ_分平台top50上新长剧播放_亿")
    concentration = read_sheet(workbook, "by集中度_top50上新长剧播放_亿")
    episodes = read_sheet(workbook, "top50上新长剧_平均集数")
    v30 = read_sheet(workbook, "top50上新长剧_集均V30_万")
    payload = {
        "unit": "亿",
        "source": "云合数据",
        "total": {
            "periods": [row["quarter"] for row in totals],
            "series": [
                {"name": "其他", "values": [clean(row.get("其他")) for row in totals]},
                {"name": "热播剧 TOP50", "values": [clean(row.get("热播剧top50")) for row in totals]},
            ],
            "totals": [clean(row.get("总计")) for row in totals],
        },
        "platform": {
            "periods": [row["quarter"] for row in platforms],
            "series": [
                {"name": name, "values": [clean(row.get(name)) for row in platforms]}
                for name in ["爱奇艺", "腾讯视频", "优酷", "芒果TV"]
            ],
        },
        "efficiency": {
            "periods": [row["quarter"] for row in concentration],
            "concentration": {
                "series": [
                    {"name": "11–50", "values": [clean(row.get("11-50")) for row in concentration]},
                    {"name": "TOP10", "values": [clean(row.get("top10")) for row in concentration]},
                ],
                "totals": [clean(row.get("top50")) for row in concentration],
            },
            "episodes": {
                "series": [
                    {"name": "TOP10", "values": [clean(row.get("top10")) for row in episodes]},
                    {"name": "11–50", "values": [clean(row.get("11-50")) for row in episodes]},
                ],
            },
            "v30": {
                "series": [
                    {"name": "TOP10", "values": [clean(row.get("top10")) for row in v30]},
                    {"name": "11–50", "values": [clean(row.get("11-50")) for row in v30]},
                ],
            },
        },
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Extracted horizontal-drama trend data to {OUTPUT}")


if __name__ == "__main__":
    main()
